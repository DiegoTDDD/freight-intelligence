from pathlib import Path

import polars as pl
from openpyxl import load_workbook

INPUT_PATH = Path("data/raw/faf5/forecast/FAF5_metadata.xlsx")
OUTPUT_PATH = Path("data/gold/faf5/commodity_lookup.parquet")


def main():
    workbook = load_workbook(
        INPUT_PATH,
        read_only=True,
        data_only=True,
    )

    sheet = workbook["Commodity (SCTG2)"]
    records = []

    for code, description in sheet.iter_rows(min_row=2, values_only=True):
        if code is None or description is None:
            continue

        records.append(
            {
                "sctg2": int(code),
                "commodity_name": str(description).strip(),
            }
        )

    workbook.close()

    lookup = pl.DataFrame(records).sort("sctg2")
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    lookup.write_parquet(OUTPUT_PATH)

    print(f"Commodity lookup created: {OUTPUT_PATH}")
    print(f"Categories: {lookup.height}")
    print(lookup)


if __name__ == "__main__":
    main()
